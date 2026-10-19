"""知识库文件文本提取：用于搜索（txt/md/docx/xlsx/xmind 提取可检索文本）。

图片、pdf、视频等无法提取文本，返回空串（后续 AI 功能再做内容理解）。
"""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path


def _read_text_file(path: Path) -> str:
    """优先 UTF-8，失败回退 GBK（常见于 Windows 生成的 txt）。"""
    raw = path.read_bytes()
    for enc in ("utf-8", "gbk", "utf-8-sig"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="ignore")


def _extract_docx(path: Path) -> str:
    """docx：python-docx 提取段落 + 表格单元格文本。"""
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    """xlsx：openpyxl 提取所有工作表单元格文本。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                vals = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if vals:
                    parts.append(" | ".join(vals))
    finally:
        wb.close()
    return "\n".join(parts)


def _extract_xmind(path: Path) -> str:
    """xmind：本质是 zip，从 content.json 提取节点标题文本（仅用于搜索）。"""
    texts: list[str] = []

    def walk(node):
        title = node.get("title")
        if title:
            texts.append(str(title))
        for child in node.get("children", {}).get("attached", []) or []:
            walk(child)

    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.endswith("content.json"):
                try:
                    data = json.loads(zf.read(name).decode("utf-8", errors="ignore"))
                except Exception:
                    continue
                for sheet in data if isinstance(data, list) else []:
                    root = sheet.get("rootTopic") or {}
                    walk(root)
    return "\n".join(texts)


def extract_file_text(path: str | Path, ext: str | None) -> str:
    """根据扩展名提取文件文本，供知识库全文检索使用。"""
    p = Path(path)
    if not p.exists():
        return ""
    ext = (ext or p.suffix.lstrip(".")).lower()
    try:
        if ext in ("txt", "md", "log", "csv"):
            return _read_text_file(p)
        if ext == "docx":
            return _extract_docx(p)
        if ext in ("xlsx", "xlsm"):
            return _extract_xlsx(p)
        if ext == "xmind":
            return _extract_xmind(p)
        # 图片/pdf/视频/zip 等：暂不提取文本
        return ""
    except Exception:
        # 提取失败不影响文档保存，仅检索不到文件内容
        return ""
